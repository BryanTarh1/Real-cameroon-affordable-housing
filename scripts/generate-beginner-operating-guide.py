from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path("/home/ubuntu/cameroon-affordable-housing")
OUTPUT = ROOT / "docs" / "AHC-Beginner-Operating-Guide.docx"

INK = "132C34"
OCHRE = "D78A1D"
PAPER = "F7F3E9"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    tc_pr.append(element)


def border_bottom(paragraph, color=OCHRE):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor.from_string(INK)
    if level == 1:
        border_bottom(paragraph)
    return paragraph


def add_text(doc, text, emphasis=None):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    if emphasis and emphasis in text:
        start, end = text.split(emphasis, 1)
        paragraph.add_run(start)
        run = paragraph.add_run(emphasis)
        run.bold = True
        paragraph.add_run(end)
    else:
        paragraph.add_run(text)
    return paragraph


def add_table(doc, headings, rows):
    table = doc.add_table(rows=1, cols=len(headings))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, heading in enumerate(headings):
        cell = table.rows[0].cells[index]
        shade(cell, INK)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    for level, size in [(1, 16), (2, 12)]:
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Affordable Housing Cameroon")
    title_run.bold = True
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(23)
    title_run.font.color.rgb = RGBColor.from_string(INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Simple Operating Guide")
    subtitle_run.font.name = "Aptos Display"
    subtitle_run.font.size = Pt(15)
    subtitle_run.font.color.rgb = RGBColor.from_string(OCHRE)
    meta = document.add_paragraph("August 2026 · For first-time visitors, Seekers, Agents, Field Moderators, and Admins")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.runs[0].font.size = Pt(9)

    notice = document.add_table(rows=1, cols=1)
    notice.alignment = WD_TABLE_ALIGNMENT.CENTER
    notice_cell = notice.cell(0, 0)
    shade(notice_cell, "F8EDD4")
    notice_cell.text = "MAIN RULE: AHC helps people find and verify rental homes. It does not hold, collect, route, guarantee, or receipt rent, deposits, viewing money, agent commissions, or tenancy settlements."
    for run in notice_cell.paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(9)

    add_heading(document, "1. Start here")
    add_text(document, "Open the AHC home page and search by city, neighbourhood, landmark, monthly rent, or your maximum Total Move-In Cash Required. This total combines the declared advance, deposit, agency fee, service fee, and first-month utilities, so you can compare homes using the full cash you need before moving.")
    add_text(document, "The map shows an approximate landmark area, not the exact compound door. Listings also show a freshness note. A listing is removed from public search if it is not reconfirmed within 14 days.")
    add_table(document, ["Your role", "Where to start", "First task"], [
        ["Seeker", "Home page", "Search homes and compare full move-in cash."],
        ["Agent", "/#/agent", "Sign in, set up a profile, and manage paid platform services."],
        ["Field Moderator", "/#/operations", "Sign in to complete assigned field verification work."],
        ["Admin", "/#/admin", "Sign in to manage AHC governance and staff controls."],
    ])

    add_heading(document, "2. How a Seeker finds a home")
    add_text(document, "Browse and filter homes without an account. When you open a property detail, sign in before using protected actions such as viewing the approved Walk-Thru, reporting an issue, requesting a viewing, or starting tracked WhatsApp contact.")
    add_text(document, "Read every item in the cost breakdown. Check the freshness note, physical-verification badge, neighbourhood badges, and the Walk-Thru if it is available. A Walk-Thru is helpful evidence, but you should still inspect a home before paying anyone.")
    add_table(document, ["Action", "Safe use"], [
        ["Compare costs", "Use the Total Move-In Cash Required, not rent alone."],
        ["Chat on WhatsApp", "Ask about that exact home. AHC records the contact intent, not the messages or a payment."],
        ["Report a problem", "Report undeclared fees, unavailable homes, or materially inaccurate information from the signed-in listing page."],
        ["Pay for a tenancy", "Deal directly with the responsible parties after your own checks. Never use AHC as a rental money intermediary."],
    ])

    add_heading(document, "3. How an Agent adds a property")
    add_text(document, "Open /#/agent and sign in with an AHC account. All suppliers use the same Agent pathway. Create a public profile with the name and WhatsApp number you want Seekers to use.")
    add_text(document, "AHC is paid from the beginning. Create the correct AHC service order, follow the official merchant instruction, and submit the provider transaction reference. An Admin must independently confirm the platform service before it becomes active. Sending a reference is not approval.")
    add_text(document, "After an Admin confirms an AHC service, open Orders and your inventory and choose Print receipt. The receipt proves only the named AHC platform service, such as Agent Access, a Listing Pass, a Featured Landmark Pin, or physical verification. It is not a rent receipt.")
    add_text(document, "Enter the real property information and every cost carefully. A listing enters independent review before it is public. After publication, use the inventory card to reconfirm availability before the 14-day freshness period ends.")

    add_heading(document, "4. How a Field Moderator works")
    add_text(document, "Open /#/operations and sign in using the account provided by AHC. Field Moderator accounts are provisioned by trusted staff, not through public registration.")
    add_text(document, "Claim suitable field work, visit the property area, and upload current evidence showing whether the home matches the listing. Record honest observations and neighbourhood details. Never invent proof, reuse old photos, or approve a home that you did not visit.")
    add_text(document, "Your commission status is shown in Operations, but it is never automatic. It remains held until required evidence is stored, any selected independent second visit is resolved, and an Admin reviews the evidence. Field Moderators do not reconcile payment references, issue receipts, handle rent, or approve payouts.")

    add_heading(document, "5. How an Admin governs the platform")
    add_text(document, "Open /#/admin and sign in. Admin access is private. The Admin workspace includes commercial settings, official service-order reconciliation, cash-flow audit, verification commission controls, trust reports, and role management.")
    add_text(document, "In the service-order queue, confirm only when provider evidence independently matches an AHC platform-service order. Confirming an order may activate the service and issue an official receipt. Reject unmatched references and write a clear reason. Never use this queue for rent, deposits, or other tenancy money.")
    add_text(document, "In the commission ledger, review evidence and any required independent audit before approving a held Field Moderator amount as payable. When reviewing trust reports, remember that a safety hold starts an investigation; it is not final proof against an Agent.")

    add_heading(document, "6. Everyday safety checklist")
    add_table(document, ["Before you act", "Safe practice"], [
        ["Before visiting", "Read the full cost breakdown and arrange contact through the listing."],
        ["Before paying tenancy money", "Visit or verify independently; AHC is not an escrow or rent-collection service."],
        ["Before accepting a receipt", "Check that it names an AHC platform service and includes an AHC receipt code."],
        ["Before uploading proof", "Use your own current visit evidence and avoid exposing unnecessary private details."],
        ["Before changing a role", "Confirm the person is trusted and the change is necessary."],
    ])

    add_heading(document, "7. If something goes wrong")
    add_text(document, "Use the complete route beginning with /#/. Use the visible Log out control before switching test roles in the same browser profile. To use more than one account on one device, use separate browser profiles or a private window.")
    add_text(document, "If an Agent cannot submit a listing, check that Agent Access and a Listing Pass are confirmed. If a Field Moderator cannot see an assignment, it may be unassigned, not route-ready, or already claimed. If an Admin cannot see a control, check the account role rather than trying to bypass the route.")
    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing_run = closing.add_run("AHC makes rental information clearer and field checks more accountable. It does not replace your own inspection, agreement, legal advice, or safe-payment judgment.")
    closing_run.italic = True
    closing_run.font.color.rgb = RGBColor.from_string(INK)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Affordable Housing Cameroon · Beginner Operating Guide · August 2026")
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("6A6358")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
